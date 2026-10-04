from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
import os


# ---------------------------------------------------------
# AISAT LOGO
# ---------------------------------------------------------

LOGO_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "assets",
    "Logo AISAT.png"
)


# ---------------------------------------------------------
# HELPER FUNCTIONS
# ---------------------------------------------------------

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()

    shd = tc_pr.find(qn("w:shd"))

    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)

    shd.set(qn("w:fill"), fill)


def set_cell_margins(
    cell,
    top=45,
    start=55,
    bottom=45,
    end=55
):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()

    tc_mar = tc_pr.first_child_found_in("w:tcMar")

    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)

    for margin, value in [
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end)
    ]:

        node = tc_mar.find(qn(f"w:{margin}"))

        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)

        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_text(
    cell,
    text,
    bold=False,
    italic=False,
    size=8.5,
    alignment=WD_ALIGN_PARAGRAPH.LEFT
):

    cell.text = ""

    paragraph = cell.paragraphs[0]

    paragraph.alignment = alignment
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0

    run = paragraph.add_run(
        str(text) if text is not None else ""
    )

    run.bold = bold
    run.italic = italic
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)

    cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )

    set_cell_margins(cell)


def set_column_width(cell, width_inches):

    cell.width = Inches(width_inches)

    tc_pr = cell._tc.get_or_add_tcPr()

    tc_w = tc_pr.find(qn("w:tcW"))

    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)

    tc_w.set(
        qn("w:w"),
        str(int(width_inches * 1440))
    )

    tc_w.set(
        qn("w:type"),
        "dxa"
    )


def set_table_borders(
    table,
    color="000000",
    size="6"
):

    tbl = table._tbl
    tbl_pr = tbl.tblPr

    borders = tbl_pr.first_child_found_in(
        "w:tblBorders"
    )

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ):

        element = borders.find(
            qn(f"w:{edge}")
        )

        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)

        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def remove_table_borders(table):

    tbl = table._tbl
    tbl_pr = tbl.tblPr

    borders = tbl_pr.first_child_found_in(
        "w:tblBorders"
    )

    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)

    for edge in (
        "top",
        "left",
        "bottom",
        "right",
        "insideH",
        "insideV"
    ):

        element = borders.find(
            qn(f"w:{edge}")
        )

        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)

        element.set(qn("w:val"), "nil")


def set_repeat_table_header(row):

    tr_pr = row._tr.get_or_add_trPr()

    tbl_header = OxmlElement("w:tblHeader")

    tbl_header.set(
        qn("w:val"),
        "true"
    )

    tr_pr.append(tbl_header)


def prevent_row_split(row):

    tr_pr = row._tr.get_or_add_trPr()

    cant_split = OxmlElement("w:cantSplit")

    tr_pr.append(cant_split)


def set_table_fixed_layout(table):

    tbl_pr = table._tbl.tblPr

    layout = tbl_pr.first_child_found_in(
        "w:tblLayout"
    )

    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)

    layout.set(
        qn("w:type"),
        "fixed"
    )


def set_row_height(
    row,
    height_twips,
    rule="atLeast"
):

    tr_pr = row._tr.get_or_add_trPr()

    tr_height = OxmlElement("w:trHeight")

    tr_height.set(
        qn("w:val"),
        str(height_twips)
    )

    tr_height.set(
        qn("w:hRule"),
        rule
    )

    tr_pr.append(tr_height)


def add_page_field(paragraph, field_name):

    run = paragraph.add_run()

    run.font.name = "Times New Roman"
    run.font.size = Pt(7.5)

    begin = OxmlElement("w:fldChar")
    begin.set(
        qn("w:fldCharType"),
        "begin"
    )

    instruction = OxmlElement("w:instrText")

    instruction.set(
        qn("xml:space"),
        "preserve"
    )

    instruction.text = field_name

    end = OxmlElement("w:fldChar")

    end.set(
        qn("w:fldCharType"),
        "end"
    )

    run._r.append(begin)
    run._r.append(instruction)
    run._r.append(end)


# ---------------------------------------------------------
# MAIN DOCUMENT GENERATOR
# ---------------------------------------------------------

def create_mom_document(
    meeting_title,
    participants,
    meeting_date,
    location,
    mom,
    photo_path=None,
    branding_path=None,
    branding_type=None,
    meeting_no="Not specified",
    absentees="Not specified",
    meeting_time="Not specified",
    mode="Not specified"
):

    document = Document()

    # -----------------------------------------------------
    # PAGE SETUP
    # -----------------------------------------------------

    section = document.sections[0]

    # A4
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)

    # Compact margins matching the reference
    section.top_margin = Inches(0.30)
    section.bottom_margin = Inches(0.30)
    section.left_margin = Inches(0.42)
    section.right_margin = Inches(0.42)

    section.header_distance = Inches(0.10)
    section.footer_distance = Inches(0.10)

    # Default font
    document.styles["Normal"].font.name = (
        "Times New Roman"
    )

    document.styles["Normal"].font.size = Pt(8.5)

    # -----------------------------------------------------
    # HEADER
    # -----------------------------------------------------

    header_table = document.add_table(
        rows=2,
        cols=1
    )

    header_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    header_table.autofit = False

    set_table_fixed_layout(header_table)
    set_table_borders(
        header_table,
        size="6"
    )

    # Logo + reference number
    header_cell = header_table.cell(0, 0)

    header_cell.text = ""

    paragraph = header_cell.paragraphs[0]

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    # Add AISAT logo
    if os.path.exists(LOGO_PATH):

        try:

            logo_run = paragraph.add_run()

            logo_run.add_picture(
                LOGO_PATH,
                width=Inches(0.65)
            )

            paragraph.add_run("   ")

        except Exception:
            pass

    # Reference number
    reference_run = paragraph.add_run(
        "AISAT/Form/QPM15/F2"
    )

    reference_run.font.name = (
        "Times New Roman"
    )

    reference_run.font.size = Pt(8)

    # Meeting Minutes
    set_cell_text(
        header_table.cell(1, 0),
        "Meeting Minutes",
        bold=True,
        size=11,
        alignment=WD_ALIGN_PARAGRAPH.CENTER
    )

    set_column_width(
        header_table.cell(0, 0),
        7.43
    )

    set_column_width(
        header_table.cell(1, 0),
        7.43
    )

    set_row_height(
        header_table.rows[0],
        260
    )

    set_row_height(
        header_table.rows[1],
        270
    )

    # Small gap
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(1)

    # -----------------------------------------------------
    # MEETING DETAILS
    # -----------------------------------------------------

    details_table = document.add_table(
        rows=5,
        cols=4
    )

    details_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    details_table.autofit = False

    set_table_fixed_layout(details_table)

    set_table_borders(
        details_table,
        size="6"
    )

    details = [

        (
            "Type",
            "Meeting",
            "Meeting No",
            meeting_no or "Not specified"
        ),

        (
            "Date",
            str(meeting_date)
            if meeting_date
            else "Not specified",

            "Venue",
            location
            if location
            else "Not specified"
        ),

        (
            "Time",
            meeting_time
            if meeting_time
            else "Not specified",

            "Mode",
            mode
            if mode
            else "Not specified"
        ),

        (
            "Attendees",
            participants
            if participants
            else "Not specified",
            "",
            ""
        ),

        (
            "Absentees",
            absentees
            if absentees
            else "Not specified",
            "",
            ""
        )
    ]

    widths = [
        0.85,
        3.20,
        1.05,
        2.33
    ]

    for row_index, row_data in enumerate(details):

        if row_index < 3:

            for col_index, value in enumerate(row_data):

                set_cell_text(
                    details_table.cell(
                        row_index,
                        col_index
                    ),
                    value,
                    bold=col_index in (0, 2),
                    size=7.8
                )

        else:

            # Merge the complete row
            details_table.cell(
                row_index,
                0
            ).merge(
                details_table.cell(
                    row_index,
                    3
                )
            )

            cell = details_table.cell(
                row_index,
                0
            )

            cell.text = ""

            p = cell.paragraphs[0]

            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)

            label_run = p.add_run(
                row_data[0] + "    "
            )

            label_run.bold = True
            label_run.font.name = (
                "Times New Roman"
            )
            label_run.font.size = Pt(7.8)

            value_run = p.add_run(
                row_data[1]
            )

            value_run.font.name = (
                "Times New Roman"
            )

            value_run.font.size = Pt(7.8)

            cell.vertical_alignment = (
                WD_CELL_VERTICAL_ALIGNMENT.CENTER
            )

            set_cell_margins(cell)

        set_row_height(
            details_table.rows[row_index],
            245
        )

    for row in details_table.rows:

        for i, width in enumerate(widths):

            if i < len(row.cells):

                set_column_width(
                    row.cells[i],
                    width
                )

    spacer = document.add_paragraph()

    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(1)

    # -----------------------------------------------------
    # AI CONTENT
    # -----------------------------------------------------

    if not isinstance(mom, dict):
        mom = {}

    agenda = mom.get(
        "agenda",
        "Not specified"
    )

    discussion_points = mom.get(
        "discussion_points",
        []
    )

    decisions = mom.get(
        "decisions",
        []
    )

    action_items = mom.get(
        "action_items",
        []
    )

    # -----------------------------------------------------
    # MAIN AISAT TABLE
    # -----------------------------------------------------

    main_table = document.add_table(
        rows=1,
        cols=3
    )

    main_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    main_table.autofit = False

    set_table_fixed_layout(main_table)

    set_table_borders(
        main_table,
        size="6"
    )

    # Header
    headers = [
        "Sl. No",
        "Agenda Item",
        "Discussions and Decisions"
    ]

    for i, header in enumerate(headers):

        set_cell_text(
            main_table.cell(0, i),
            header,
            bold=True,
            size=7.8,
            alignment=WD_ALIGN_PARAGRAPH.CENTER
        )

    set_repeat_table_header(
        main_table.rows[0]
    )

    set_row_height(
        main_table.rows[0],
        270
    )

    # -----------------------------------------------------
    # OFFICIAL TEMPLATE ROWS
    # -----------------------------------------------------

    rows = []

    # 1
    rows.append(
        (
            "Prayer",
            ""
        )
    )

    # 2
    rows.append(
        (
            "Welcome",
            ""
        )
    )

    # 3
    rows.append(
        (
            "Previous Minutes Approval",
            "Update Action Taken Report at the end"
        )
    )

    # 4
    rows.append(
        (
            "General Update of activities",
            ""
        )
    )

    # AI agenda
    if agenda and agenda != "Not specified":

        rows.append(
            (
                "Agenda",
                agenda
            )
        )

    # Discussions
    for point in discussion_points:

        if str(point).strip():

            rows.append(
                (
                    "Discussion",
                    str(point).strip()
                )
            )

    # Decisions
    for decision in decisions:

        if str(decision).strip():

            rows.append(
                (
                    "Decision",
                    str(decision).strip()
                )
            )

    # Action items
    for action in action_items:

        if not isinstance(action, dict):
            continue

        task = str(
            action.get(
                "task",
                "Not specified"
            )
        ).strip()

        responsible = str(
            action.get(
                "responsible",
                "Not specified"
            )
        ).strip()

        if not task:
            task = "Not specified"

        if not responsible:
            responsible = "Not specified"

        rows.append(
            (
                "Action Item",
                f"{task} — Responsibility: {responsible}"
            )
        )

    # Varia
    rows.append(
        (
            "Varia",
            ""
        )
    )

    # -----------------------------------------------------
    # ADD CONTENT ROWS
    # -----------------------------------------------------

    for index, row_data in enumerate(
        rows,
        start=1
    ):

        agenda_item = row_data[0]
        content = row_data[1]

        row = main_table.add_row()

        prevent_row_split(row)

        set_cell_text(
            row.cells[0],
            str(index),
            size=7.5,
            alignment=WD_ALIGN_PARAGRAPH.CENTER
        )

        set_cell_text(
            row.cells[1],
            agenda_item,
            bold=agenda_item in (
                "Agenda",
                "Decision",
                "Action Item"
            ),
            size=7.5
        )

        set_cell_text(
            row.cells[2],
            content,
            size=7.5
        )

        set_row_height(
            row,
            235
        )

    # -----------------------------------------------------
    # COLUMN WIDTHS
    # -----------------------------------------------------

    for row in main_table.rows:

        set_column_width(
            row.cells[0],
            0.45
        )

        set_column_width(
            row.cells[1],
            2.05
        )

        set_column_width(
            row.cells[2],
            4.93
        )

    # -----------------------------------------------------
    # PHOTO OF MEETING
    # -----------------------------------------------------

    photo_row = main_table.add_row()

    prevent_row_split(photo_row)

    photo_number = len(rows) + 1

    set_cell_text(
        photo_row.cells[0],
        str(photo_number),
        size=7.5,
        alignment=WD_ALIGN_PARAGRAPH.CENTER
    )

    set_cell_text(
        photo_row.cells[1],
        "Photo of the meeting",
        bold=True,
        size=7.5
    )

    photo_cell = photo_row.cells[2]

    photo_cell.text = ""

    paragraph = photo_cell.paragraphs[0]

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.CENTER
    )

    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    if photo_path and os.path.exists(photo_path):

        try:

            run = paragraph.add_run()

            run.add_picture(
                photo_path,
                width=Inches(2.20)
            )

        except Exception:

            run = paragraph.add_run(
                "Geotagged Photo"
            )

            run.italic = True
            run.font.name = (
                "Times New Roman"
            )
            run.font.size = Pt(7.5)

    else:

        run = paragraph.add_run(
            "Geotagged Photo"
        )

        run.italic = True
        run.font.name = (
            "Times New Roman"
        )
        run.font.size = Pt(7.5)

    photo_cell.vertical_alignment = (
        WD_CELL_VERTICAL_ALIGNMENT.CENTER
    )

    set_cell_margins(
        photo_cell,
        top=35,
        bottom=35
    )

    for row in [photo_row]:

        set_column_width(
            row.cells[0],
            0.45
        )

        set_column_width(
            row.cells[1],
            2.05
        )

        set_column_width(
            row.cells[2],
            4.93
        )

    # -----------------------------------------------------
    # ACTION TAKEN REPORT
    # -----------------------------------------------------

    atr_table = document.add_table(
        rows=2,
        cols=4
    )

    atr_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    atr_table.autofit = False

    set_table_fixed_layout(atr_table)

    set_table_borders(
        atr_table,
        size="6"
    )

    # Title
    title_cell = atr_table.cell(0, 0)

    for i in range(1, 4):

        title_cell = title_cell.merge(
            atr_table.cell(0, i)
        )

    set_cell_text(
        title_cell,
        "Action Taken Report",
        bold=True,
        size=8.5,
        alignment=WD_ALIGN_PARAGRAPH.CENTER
    )

    # Header
    atr_headers = [
        "Reference",
        "Action Item",
        "Status",
        "Responsibility"
    ]

    for i, header in enumerate(
        atr_headers
    ):

        set_cell_text(
            atr_table.cell(1, i),
            header,
            bold=True,
            size=7.5,
            alignment=WD_ALIGN_PARAGRAPH.CENTER
        )

    set_row_height(
        atr_table.rows[0],
        260
    )

    set_row_height(
        atr_table.rows[1],
        250
    )

    # Action rows
    if action_items:

        for action in action_items:

            if not isinstance(action, dict):
                continue

            task = str(
                action.get(
                    "task",
                    "Not specified"
                )
            ).strip()

            responsible = str(
                action.get(
                    "responsible",
                    "Not specified"
                )
            ).strip()

            if not task:
                task = "Not specified"

            if not responsible:
                responsible = "Not specified"

            row = atr_table.add_row()

            prevent_row_split(row)

            values = [
                "MeetingNo.ItemSl.No",
                task,
                "Not specified",
                responsible
            ]

            for i, value in enumerate(values):

                set_cell_text(
                    row.cells[i],
                    value,
                    italic=(i == 0),
                    size=7,
                    alignment=(
                        WD_ALIGN_PARAGRAPH.CENTER
                        if i in (0, 2)
                        else WD_ALIGN_PARAGRAPH.LEFT
                    )
                )

            set_row_height(
                row,
                230
            )

    else:

        row = atr_table.add_row()

        values = [
            "MeetingNo.ItemSl.No",
            "No specific action items identified.",
            "Not specified",
            "Not specified"
        ]

        for i, value in enumerate(values):

            set_cell_text(
                row.cells[i],
                value,
                italic=(i == 0),
                size=7,
                alignment=(
                    WD_ALIGN_PARAGRAPH.CENTER
                    if i in (0, 2)
                    else WD_ALIGN_PARAGRAPH.LEFT
                )
            )

        set_row_height(
            row,
            230
        )

    # ATR widths
    for row in atr_table.rows:

        set_column_width(
            row.cells[0],
            1.18
        )

        set_column_width(
            row.cells[1],
            3.55
        )

        set_column_width(
            row.cells[2],
            0.85
        )

        set_column_width(
            row.cells[3],
            1.85
        )

    # -----------------------------------------------------
    # SIGNATURES
    # -----------------------------------------------------

    signature_table = document.add_table(
        rows=2,
        cols=3
    )

    signature_table.alignment = (
        WD_TABLE_ALIGNMENT.CENTER
    )

    signature_table.autofit = False

    set_table_fixed_layout(
        signature_table
    )

    remove_table_borders(
        signature_table
    )

    roles = [
        "Secretary",
        "Coordinator",
        "Chairman"
    ]

    for i, role in enumerate(roles):

        set_cell_text(
            signature_table.cell(0, i),
            "Name and Signature",
            size=7.5,
            alignment=WD_ALIGN_PARAGRAPH.CENTER
        )

        set_cell_text(
            signature_table.cell(1, i),
            role,
            size=7.5,
            alignment=WD_ALIGN_PARAGRAPH.CENTER
        )

        set_column_width(
            signature_table.cell(0, i),
            2.48
        )

        set_column_width(
            signature_table.cell(1, i),
            2.48
        )

    set_row_height(
        signature_table.rows[0],
        330
    )

    set_row_height(
        signature_table.rows[1],
        230
    )

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    footer = section.footer

    paragraph = footer.paragraphs[0]

    paragraph.alignment = (
        WD_ALIGN_PARAGRAPH.RIGHT
    )

    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run(
        "Page "
    )

    run.font.name = (
        "Times New Roman"
    )

    run.font.size = Pt(7.5)

    add_page_field(
        paragraph,
        "PAGE"
    )

    run = paragraph.add_run(
        "/"
    )

    run.font.name = (
        "Times New Roman"
    )

    run.font.size = Pt(7.5)

    add_page_field(
        paragraph,
        "NUMPAGES"
    )

    return document